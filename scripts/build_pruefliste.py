"""Erzeugt output/Pruefliste_Fachschaften.xlsx: je Fach ein Blatt mit allem, was fachlich zu prüfen ist.

Themenvorschläge und SDG-Zuordnungen sind eigene Bündelungen, Schwerpunkte der neu erfassten Fächer
wurden automatisch ausgelesen oder von Hand zusammengefasst. Die Fachschaften bewerten jede Zeile.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output", "Pruefliste_Fachschaften.xlsx")
PILOT_GY = {"Biologie", "Chemie", "Physik", "Erdkunde", "Politik und Wirtschaft"}
SFN = {"hs": "HS", "rs": "RS", "gy": "GY"}


def lies(name):
    return list(csv.DictReader(open(os.path.join(ROOT, "data", name), encoding="utf-8")))


F = "Arial"
HFILL, HFONT = PatternFill("solid", fgColor="1F4E5F"), Font(name=F, bold=True, color="FFFFFF", size=10)
TYPFILL = {"Themenvorschlag": "FFF4D6", "SDG-Zuordnung": "E3F4E8", "Schwerpunkt": "EEF3F8"}
faecher = [r["fach"] for r in lies("faecher.csv")]
sdg = {r["nr"]: r["titel"] for r in lies("bne_sdg.csv")}

zeilen = {f: [] for f in faecher}
for r in lies("themen.csv"):
    zeilen[r["fach"]].append(("Themenvorschlag", r["thema"], r["inhaltsfeld"], r["stichwort"], "", r["stufe"] or r["seite"]))
for r in lies("bne_vorschlaege.csv"):
    zeilen[r["fach"]].append(("SDG-Zuordnung", f"SDG {r['sdg']}: {sdg[r['sdg']]}", r["inhaltsfeld"], r["stichwort"], "", ""))
# Schwerpunkte: gleiche Formulierung in mehreren Fassungen zusammenfassen
sp = {}
for r in lies("schwerpunkte_sek1.csv"):
    if r["fach"] in PILOT_GY:
        continue
    k = (r["fach"], r["inhaltsfeld"], r["schwerpunkt"], r["stufe_typ"])
    e = sp.setdefault(k, {"sf": [], "seiten": [], "stufe": r["stufe"], "hin": r["hinweis"]})
    e["sf"].append(SFN[r["schulform"]]); e["seiten"].append(f'{SFN[r["schulform"]]} {r["seite"]}')
for (fach, feld, text, typ), e in sp.items():
    zeilen[fach].append((f"Schwerpunkt", e["stufe"], feld, text, ", ".join(e["sf"]), "; ".join(e["seiten"]) + (f" ({e['hin']})" if e["hin"] else "")))

wb = Workbook()
ov = wb.active
ov.title = "Anleitung"
text = [
    ("Prüfliste für die Fachschaften", Font(name=F, bold=True, size=14)),
    ("Kompetenzgraph Hessen, Sekundarstufe I. Stand 27.10.2026.", Font(name=F, italic=True, size=10)),
    ("", None),
    ("Worum es geht", Font(name=F, bold=True, size=11)),
    ("Jedes Blatt enthält für ein Fach die Einträge, die fachlich geprüft werden sollen:", Font(name=F, size=10)),
    ("• Themenvorschläge: eigene Bündelung über gemeinsame Stichworte. Sie stehen so nicht im Kerncurriculum.", Font(name=F, size=10)),
    ("• SDG-Zuordnungen: Verbindung von Inhaltsfeldern zu den 17 Nachhaltigkeitszielen. Die Kerncurricula nennen keine SDGs.", Font(name=F, size=10)),
    ("• Schwerpunkte: aus dem Kerncurriculum ausgelesen (automatisch oder von Hand zusammengefasst). Bitte Wortlaut und Seitenangabe am Original prüfen.", Font(name=F, size=10)),
    ("", None),
    ("So geht es", Font(name=F, bold=True, size=11)),
    ("1. In der Spalte „Bewertung“ aus der Liste wählen: passt, passt teilweise, passt nicht.", Font(name=F, size=10)),
    ("2. Bei „passt teilweise“ oder „passt nicht“ in „Anmerkung“ kurz notieren, was geändert werden soll.", Font(name=F, size=10)),
    ("3. Name der Lehrerin oder des Lehrers und Datum eintragen. Die Datei zurückgeben, die Änderungen werden in die Datenbasis übernommen.", Font(name=F, size=10)),
    ("", None),
    ("Übersicht", Font(name=F, bold=True, size=11)),
]
for t, font in text:
    ov.append([t])
    if font:
        ov.cell(ov.max_row, 1).font = font
ov.append(["Fach", "Themenvorschläge", "SDG-Zuordnungen", "Schwerpunkte", "Summe"])
for c in ov[ov.max_row]:
    c.fill, c.font = HFILL, HFONT
ov.column_dimensions["A"].width = 26
for L in "BCDE":
    ov.column_dimensions[L].width = 18

KOPF = ["Typ", "Thema / SDG / Stufe", "Inhaltsfeld", "Stichwort bzw. Schwerpunkt", "KC-Fassung", "Seite bzw. Stufe laut KC", "Bewertung", "Anmerkung", "Geprüft von", "Datum"]
BREITE = [16, 30, 36, 60, 11, 30, 16, 40, 18, 12]
for fach in faecher:
    rows = zeilen[fach]
    name = {"Politik und Wirtschaft": "PoWi", "Moderne Fremdsprachen": "Fremdsprachen", "Evangelische Religion": "Ev. Religion", "Katholische Religion": "Kath. Religion"}.get(fach, fach)
    ws = wb.create_sheet(name)
    ws.append(KOPF)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HFILL, HFONT, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r) + ["", "", "", ""])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = Font(name=F, size=10), Alignment(wrap_text=True, vertical="top")
        row[0].fill = PatternFill("solid", fgColor=TYPFILL[row[0].value])
    for i, w in enumerate(BREITE, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if rows:
        dv = DataValidation(type="list", formula1='"passt,passt teilweise,passt nicht"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"G2:G{len(rows) + 1}")
    n = lambda t: sum(1 for r in rows if r[0] == t)
    ov.append([fach, n("Themenvorschlag"), n("SDG-Zuordnung"), n("Schwerpunkt"), len(rows)])
    for c in ov[ov.max_row]:
        c.font = Font(name=F, size=10)

wb.save(OUT)
print(OUT, sum(len(v) for v in zeilen.values()), "Zeilen")
