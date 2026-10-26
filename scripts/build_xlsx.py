import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "data"))
from kerncurricula import *

OUT = os.path.join(ROOT, "output", "Kerncurricula_Hessen_Sek1.xlsx")


def lies(name):
    return list(csv.DictReader(open(os.path.join(ROOT, "data", name), encoding="utf-8")))



FAECHER = [(r["key"], r["fach"], r["kurz"], r["farbe"]) for r in lies("faecher.csv")]  # wie in der Webseite

# ---------------- Workbook ----------------
F = "Arial"
HFILL = PatternFill("solid", fgColor="1F4E5F")
HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BFONT = Font(name=F, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D7DE")
BORDER = Border(bottom=thin)
def hell(c):  # Fachfarbe stark aufgehellt als Zellhintergrund
    return "".join(f"{int(c[i:i + 2], 16) + (255 - int(c[i:i + 2], 16)) * 4 // 5:02X}" for i in (1, 3, 5))
FACHFILL = {f: hell(c) for _, f, _, c in FAECHER}

def sheet(ws, headers, rows, widths, fach_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HFILL, HFONT, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BFONT, WRAP, BORDER
        if fach_col is not None:
            fach = row[fach_col].value
            if fach in FACHFILL:
                row[fach_col].fill = PatternFill("solid", fgColor=FACHFILL[fach])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

wb = Workbook()
ov = wb.active
ov.title = "Übersicht"

# IDs
IFP = {(fach, feld): bk for fach, bk, feld, seite, hin in IF}
if_rows = [(f"IF-{i:03d}", r["fach"], r["bereich"] or IFP.get((r["fach"], r["inhaltsfeld"]), ""), r["inhaltsfeld"],
            r["seite_hs"], r["seite_rs"], r["seite_gy"], r["hinweis"]) for i, r in enumerate(lies("inhaltsfelder_sek1.csv"), 1)]
SFN = {"hs": "Hauptschule", "rs": "Realschule", "gy": "Gymnasium"}
sp_rows = [r[:2] + ("Gymnasium",) + r[2:] + ("",) for r in SP]
sp_rows += [r[:2] + ("Hauptschule, Realschule", r[2], r[3], r[4], "Kap. 6 lautet in allen drei Fassungen gleich")
            for r in SP if r[0] in ("Erdkunde", "Politik und Wirtschaft")]
sp_rows += [(r["fach"], r["inhaltsfeld"], SFN[r["schulform"]], r["stufe"], r["schwerpunkt"], r["seite"], r["hinweis"])
            for r in lies("schwerpunkte_sek1.csv")]
sp_rows = [(f"SP-{i:03d}",) + r for i, r in enumerate(sp_rows, 1)]
SDG = {r["nr"]: r["titel"] for r in lies("bne_sdg.csv")}
bne_rows = [(r["fach"], r["inhaltsfeld"] or "(fachweiter Bildungsstandard)", r["seite_hs"], r["seite_rs"], r["seite_gy"], r["zitat"])
            for r in lies("bne_belege.csv")]
sdg_rows = [(int(r["sdg"]), SDG[r["sdg"]], r["fach"], r["inhaltsfeld"], r["stichwort"], "Vorschlag, fachlich zu prüfen")
            for r in lies("bne_vorschlaege.csv")]
ak_rows = [(r["thema"], r["fach"], r["inhaltsfeld"], r["stichwort"], r["stufe"] or r["seite"], "Vorschlag, fachlich zu prüfen")
           for r in lies("themen.csv")]

ws = wb.create_sheet("Inhaltsfelder")
sheet(ws, ["ID", "Fach", "Einordnung im KC", "Inhaltsfeld", "Seite HS (Kap. 6)", "Seite RS (Kap. 6)", "Seite GY (Kap. 6)", "Hinweis"],
      if_rows, [9, 20, 36, 52, 11, 11, 11, 40], fach_col=1)
ws = wb.create_sheet("Schwerpunkte")
sheet(ws, ["ID", "Fach", "Inhaltsfeld", "KC-Fassung", "Stufe laut KC", "Inhaltlicher Schwerpunkt", "Seite", "Hinweis"],
      sp_rows, [9, 20, 40, 16, 34, 70, 10, 34], fach_col=1)
ws = wb.create_sheet("Kompetenzbereiche")
sheet(ws, ["Fach", "Kompetenzbereich", "Teilbereiche / Operatoren", "Seite", "Hinweis"],
      KB, [20, 34, 80, 8, 36], fach_col=0)
ws = wb.create_sheet("Querbezüge belegt")
sheet(ws, ["Fach (Quelle)", "Bezug zu", "Aussage im KC (sinngemäß)", "Fundstelle"],
      QB, [22, 28, 90, 26])
ws = wb.create_sheet("Anknüpfungspunkte")
sheet(ws, ["Thema", "Fach", "Inhaltsfeld", "Stichwort aus dem KC", "Stufe laut KC bzw. Seite (Kap. 6)", "Status"],
      ak_rows, [26, 20, 40, 60, 30, 26], fach_col=1)
ws = wb.create_sheet("BNE Belege")
sheet(ws, ["Fach", "Inhaltsfeld", "Seite HS", "Seite RS", "Seite GY", "Aussage im KC (gekürzt)"],
      bne_rows, [20, 44, 10, 10, 10, 90], fach_col=0)
ws = wb.create_sheet("SDG-Vorschläge")
sheet(ws, ["SDG", "Nachhaltigkeitsziel", "Fach", "Inhaltsfeld", "Stichwort", "Status"],
      sdg_rows, [7, 34, 20, 48, 40, 26], fach_col=2)
ws = wb.create_sheet("Hinweise und Lücken")
sheet(ws, ["Punkt", "Hinweis"], HIN, [20, 110])

# Übersicht
ov["A1"] = "Kerncurricula Hessen, Sek I (Haupt-, Real- und Gymnasialfassung) mit BNE"
ov["A1"].font = Font(name=F, bold=True, size=14)
ov["A2"] = "Stand der Auswertung: 26.10.2026. Quelle: Hessisches Kultusministerium, Seite 'Sekundarstufe I – Kerncurricula'."
ov["A2"].font = Font(name=F, size=10, italic=True)
ov.append([])
ov.append(["Fach", "Inhaltsfelder", "Schwerpunkte", "Anknüpfungspunkte", "SDG-Vorschläge", "PDF Hauptschule", "PDF Realschule", "PDF Gymnasium"])
for c in ov[4]:
    c.fill, c.font = HFILL, HFONT
URLS = {}
for q in lies("quellen_sek1.csv"):
    if q["dokument"] == "Kerncurriculum":
        URLS[(q["fach"], {"Hauptschule": 6, "Realschule": 7}.get(q["schulform"], 8))] = q["url"]
r = 5
for _, fach, _, _ in FAECHER:
    ov.cell(r, 1, fach)
    ov.cell(r, 2, f'=COUNTIF(Inhaltsfelder!B:B,A{r})')
    ov.cell(r, 3, f'=COUNTIF(Schwerpunkte!B:B,A{r})')
    ov.cell(r, 4, f'=COUNTIF(Anknüpfungspunkte!B:B,A{r})')
    ov.cell(r, 5, f"=COUNTIF('SDG-Vorschläge'!C:C,A{r})")
    for col in range(1, 6):
        ov.cell(r, col).font = BFONT
    for col in (6, 7, 8):
        u = URLS.get((fach, col))
        if u:
            ov.cell(r, col, "PDF").hyperlink = u
            ov.cell(r, col).font = Font(name=F, size=10, color="0563C1", underline="single")
    ov.cell(r, 1).fill = PatternFill("solid", fgColor=FACHFILL[fach])
    r += 1
ov.cell(r, 1, "Summe").font = Font(name=F, bold=True, size=10)
for col, L in [(2, "B"), (3, "C"), (4, "D"), (5, "E")]:
    ov.cell(r, col, f"=SUM({L}5:{L}{r-1})").font = Font(name=F, bold=True, size=10)
r += 2
ov.cell(r, 1, "Tabellenblätter").font = Font(name=F, bold=True, size=11)
for name, desc in [
    ("Inhaltsfelder", "Alle Inhaltsfelder je Fach mit Einordnung und Seitenangabe je KC-Fassung (leer = in dieser Fassung nicht vorhanden)"),
    ("Schwerpunkte", "Inhaltliche Schwerpunkte je Inhaltsfeld, KC-Fassung und Stufe, bisher für die fünf Pilotfächer"),
    ("Kompetenzbereiche", "Fachliche Kompetenzbereiche und Teilbereiche"),
    ("Querbezüge belegt", "Fächerübergreifende Bezüge, die ausdrücklich im KC stehen"),
    ("Anknüpfungspunkte", "Themenvorschläge über gemeinsame Stichworte, zu prüfen"),
    ("BNE Belege", "Stellen, an denen das KC selbst Nachhaltigkeit nennt (BNE ist Aufgabengebiet nach § 6 Abs. 4 HSchG)"),
    ("SDG-Vorschläge", "Zuordnung von Inhaltsfeldern zu den 17 Nachhaltigkeitszielen, zu prüfen"),
    ("Hinweise und Lücken", "Was fehlt oder abweicht"),
]:
    r += 1
    ov.cell(r, 1, name).font = Font(name=F, bold=True, size=10)
    ov.cell(r, 2, desc).font = BFONT
ov.column_dimensions["A"].width = 24
ov.column_dimensions["B"].width = 16
ov.column_dimensions["C"].width = 16
for L, w in zip("DEFGH", (20, 16, 16, 16, 16)):
    ov.column_dimensions[L].width = w

wb.save(OUT)
print(OUT, len(if_rows), len(sp_rows), len(KB), len(QB), len(ak_rows), len(bne_rows), len(sdg_rows))
